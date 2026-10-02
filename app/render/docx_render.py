"""Blocks -> DOCX via python-docx.

Same block model as the PDF renderer. Change diffs map to native Word
strikethrough (deletions) and text highlight (insertions), so reviewers
can comment and track changes in Word.
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.docmodel.blocks import (
    AnyBlock, Bullets, DiffPara, H1, H2, H3, KVTable, Label, Note,
    Numbered, PageBreak, Paragraph, PlanDocument, Table,
)

BADGE_COLORS = {
    "IMPROVE": RGBColor(0x1D, 0x4E, 0xD8), "MERGE": RGBColor(0x7C, 0x3A, 0xED),
    "REMOVE": RGBColor(0xB9, 0x1C, 0x1C), "KEEP": RGBColor(0x15, 0x80, 0x3D),
    "ADD": RGBColor(0x15, 0x80, 0x3D), "NEW": RGBColor(0x15, 0x80, 0x3D),
}
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB9, 0x1C, 0x1C)


def _field(paragraph, code: str) -> None:
    """Insert a Word field (PAGE / NUMPAGES) into a paragraph."""
    run = paragraph.add_run()
    for el, attr, val in (
        ("w:fldChar", "w:fldCharType", "begin"),
        ("w:instrText", "xml:space", "preserve"),
        ("w:fldChar", "w:fldCharType", "end"),
    ):
        node = OxmlElement(el)
        if el == "w:instrText":
            node.text = code
        else:
            node.set(qn(attr), val)
        run._r.append(node)


def _page_footer(doc: Document) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Page ")
        run.font.size = Pt(8)
        _field(p, "PAGE")
        mid = p.add_run(" of ")
        mid.font.size = Pt(8)
        _field(p, "NUMPAGES")
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = GREY


def _heading(doc: Document, text: str, size: int, badge: str | None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if size >= 13 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    if size >= 13:  # H2 gets a bottom rule like the PDF
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:color"), "111111")
        pBdr.append(bottom)
        pPr.append(pBdr)
    if badge:
        b = p.add_run(f"[{badge.upper()}] ")
        b.bold = True
        b.font.size = Pt(max(size - 5, 7))
        b.font.color.rgb = BADGE_COLORS.get(badge.upper(), GREY)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)


def _add_table(doc: Document, header: list[str], rows: list[list[str]],
               first_col_bold: bool = False) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8.5)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run()
            _add_multiline(run, val)
            run.font.size = Pt(8.5)
            if first_col_bold and j == 0:
                run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_multiline(run, text: str) -> None:
    """Write text into a run, converting newlines to Word line breaks."""
    parts = text.split("\n")
    for i, part in enumerate(parts):
        if i:
            run.add_break()
        run.add_text(part)


def _render_block(doc: Document, b: AnyBlock) -> None:
    if isinstance(b, H1):
        p = doc.add_paragraph()
        r = p.add_run(b.text)
        r.bold = True
        r.font.size = Pt(20)
    elif isinstance(b, H2):
        _heading(doc, b.text, 13, b.badge)
    elif isinstance(b, H3):
        _heading(doc, b.text, 11, b.badge)
    elif isinstance(b, Paragraph):
        p = doc.add_paragraph(b.text)
        p.paragraph_format.space_after = Pt(4)
    elif isinstance(b, Note):
        p = doc.add_paragraph()
        r = p.add_run(b.text)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = GREY
    elif isinstance(b, Label):
        p = doc.add_paragraph()
        r = p.add_run(b.text)
        r.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = GREY
    elif isinstance(b, PageBreak):
        doc.add_page_break()
    elif isinstance(b, KVTable):
        _add_table(doc, ["", ""], b.pairs, first_col_bold=True)
    elif isinstance(b, Table):
        if b.caption:
            p = doc.add_paragraph()
            r = p.add_run(b.caption)
            r.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = GREY
        _add_table(doc, b.header, b.rows)
    elif isinstance(b, Bullets):
        for item in b.items:
            doc.add_paragraph(item, style="List Bullet")
    elif isinstance(b, Numbered):
        for item in b.items:
            doc.add_paragraph(item, style="List Number")
    elif isinstance(b, DiffPara):
        p = doc.add_paragraph()
        for run in b.runs:
            r = p.add_run(run.text)
            if run.style == "del":
                r.font.strike = True
                r.font.color.rgb = RED
            elif run.style == "ins":
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    else:
        raise ValueError(f"unknown block: {type(b)}")


def render_docx(doc: PlanDocument, out_path: str) -> None:
    d = Document()
    for section in d.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)

    t = d.add_paragraph()
    r = t.add_run(doc.title)
    r.bold = True
    r.font.size = Pt(22)
    if doc.subtitle:
        s = d.add_paragraph()
        r = s.add_run(doc.subtitle)
        r.font.size = Pt(14)
    if doc.market_note:
        n = d.add_paragraph()
        r = n.add_run(doc.market_note)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY

    for b in doc.blocks:
        _render_block(d, b)

    _page_footer(d)
    d.save(out_path)
